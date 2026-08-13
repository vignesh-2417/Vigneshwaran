#!/usr/bin/env python3
"""Generate the Vulcan Industries Salesforce User Guide (.docx)."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
MEDIA = ROOT / "media"
OUTPUT = ROOT / "Vulcan-Industries-Salesforce-User-Guide.docx"

# Salesforce-inspired professional palette
SF_BLUE = RGBColor(0x01, 0x7E, 0xC5)
DARK = RGBColor(0x18, 0x18, 0x18)
GRAY = RGBColor(0x5C, 0x5C, 0x5C)
CALLOUT_BG = "EEF4FF"
TIP_BG = "F3F8F0"
WARN_BG = "FFF4E5"
IMPORTANT_BG = "FCE8E8"


def set_run_font(run, *, size=11, bold=False, color=DARK, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def set_cell_shading(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_paragraph_spacing(paragraph, before=0, after=8, line=1.15):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line


def add_horizontal_line(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "017EC5")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_heading_styled(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = SF_BLUE if level <= 2 else DARK
        run.font.name = "Calibri"
    set_paragraph_spacing(h, before=16 if level == 1 else 12, after=6)
    return h


def add_body(doc, text, *, bold=False, italic=False):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=8)
    run = p.add_run(text)
    set_run_font(run, bold=bold)
    run.italic = italic
    return p


def add_rich_paragraph(doc, parts, *, after=8):
    """parts: list of (text, bold, italic) tuples."""
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=0, after=after)
    for text, bold, italic in parts:
        run = p.add_run(text)
        set_run_font(run, bold=bold)
        run.italic = italic
    return p


def add_bullet(doc, text, *, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    set_paragraph_spacing(p, before=0, after=4)
    p.paragraph_format.left_indent = Inches(0.25 + 0.25 * level)
    # Support simple **bold** markers via segments
    segments = []
    remaining = text
    while remaining:
        if "**" in remaining:
            before, rest = remaining.split("**", 1)
            if before:
                segments.append((before, False))
            if "**" in rest:
                bold_text, remaining = rest.split("**", 1)
                segments.append((bold_text, True))
            else:
                segments.append((rest, False))
                remaining = ""
        else:
            segments.append((remaining, False))
            remaining = ""
    for t, bold in segments:
        run = p.add_run(t)
        set_run_font(run, bold=bold)
    return p


def add_numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.clear()
    set_paragraph_spacing(p, before=0, after=4)
    segments = []
    remaining = text
    while remaining:
        if "**" in remaining:
            before, rest = remaining.split("**", 1)
            if before:
                segments.append((before, False))
            if "**" in rest:
                bold_text, remaining = rest.split("**", 1)
                segments.append((bold_text, True))
            else:
                segments.append((rest, False))
                remaining = ""
        else:
            segments.append((remaining, False))
            remaining = ""
    for t, bold in segments:
        run = p.add_run(t)
        set_run_font(run, bold=bold)
    return p


def add_callout(doc, title, body, bg=CALLOUT_BG):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, bg)

    # Clear default paragraph and write title + body
    cell.paragraphs[0].clear()
    title_run = cell.paragraphs[0].add_run(title)
    set_run_font(title_run, size=10, bold=True, color=SF_BLUE)
    set_paragraph_spacing(cell.paragraphs[0], before=4, after=2)

    body_p = cell.add_paragraph()
    body_run = body_p.add_run(body)
    set_run_font(body_run, size=10, color=DARK)
    set_paragraph_spacing(body_p, before=0, after=4)

    # Add spacing after table
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=0, after=8)
    return table


def add_figure(doc, image_name, caption, width=6.0):
    path = MEDIA / image_name
    if not path.exists():
        add_body(doc, f"[Missing image: {image_name}]", italic=True)
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(p, before=8, after=4)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(cap, before=0, after=12)
    run = cap.add_run(caption)
    set_run_font(run, size=9, color=GRAY)
    run.italic = True


def add_kv_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(rows):
        c0, c1 = table.cell(i, 0), table.cell(i, 1)
        c0.text = ""
        c1.text = ""
        r0 = c0.paragraphs[0].add_run(k)
        set_run_font(r0, size=10, bold=True)
        r1 = c1.paragraphs[0].add_run(v)
        set_run_font(r1, size=10)
        if i == 0:
            set_cell_shading(c0, "017EC5")
            set_cell_shading(c1, "017EC5")
            r0.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            r1.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        else:
            set_cell_shading(c0, "F4F6F9")
    spacer = doc.add_paragraph()
    set_paragraph_spacing(spacer, before=0, after=10)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK

    for i in range(1, 4):
        style = styles[f"Heading {i}"]
        style.font.name = "Calibri"
        style.font.color.rgb = SF_BLUE
        style.font.bold = True
        if i == 1:
            style.font.size = Pt(18)
        elif i == 2:
            style.font.size = Pt(14)
        else:
            style.font.size = Pt(12)


def build():
    doc = Document()
    configure_styles(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    # ----- Title page -----
    for _ in range(2):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("VULCAN INDUSTRIES")
    set_run_font(run, size=28, bold=True, color=SF_BLUE, name="Calibri")
    set_paragraph_spacing(title, before=36, after=6)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Salesforce User Guide")
    set_run_font(run, size=22, bold=True, color=DARK)
    set_paragraph_spacing(subtitle, before=0, after=6)

    tagline = doc.add_paragraph()
    tagline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = tagline.add_run("Quote-to-Order Process")
    set_run_font(run, size=14, color=GRAY)
    set_paragraph_spacing(tagline, before=0, after=18)
    add_horizontal_line(tagline)

    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = intro.add_run(
        "End-user documentation for Account, Address, Contract & Pricing,\n"
        "Quote (STLE), Approvals, Order, and SyteLine integration."
    )
    set_run_font(run, size=11, color=GRAY)
    set_paragraph_spacing(intro, before=18, after=24)

    add_kv_table(
        doc,
        [
            ("Document Field", "Value"),
            ("Application", "Vulcan Industries / Vulcan Industrial"),
            ("Audience", "Sales and Customer Service end users"),
            ("Scope", "Account → Address → Contract → Quote → Approval → Order → SyteLine"),
            ("Document Type", "User Guide"),
            ("Version", "1.0"),
        ],
    )

    doc.add_page_break()

    # ----- TOC -----
    add_heading_styled(doc, "Table of Contents", level=1)
    toc_items = [
        "1. Overview",
        "2. Create an Account",
        "3. Create an Address for the Account",
        "4. Contract and Contract Pricing",
        "5. Create a Quote (from Account)",
        "6. Add Products in STLE",
        "7. Generate and Send the Quote Document",
        "8. Quote Approvals",
        "9. Create an Order from a Quote",
        "10. Order Documents and Activation",
        "11. Quick Reference — Key Tips & Rules",
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        set_paragraph_spacing(p, before=2, after=2)
        run = p.add_run(item)
        set_run_font(run, size=11, color=DARK)

    doc.add_page_break()

    # ----- 1. Overview -----
    add_heading_styled(doc, "1. Overview", level=1)
    add_body(
        doc,
        "This guide describes the end-to-end Quote-to-Order process in Salesforce "
        "for Vulcan Industries. Follow the sections in order for the standard path, "
        "or jump to a specific topic using the table of contents.",
    )
    add_body(doc, "Process flow:", bold=True)
    add_body(
        doc,
        "Account  →  Address  →  Contract + Contract Pricing  →  Quote (STLE)  →  Approval  →  Order  →  SyteLine",
        italic=True,
    )
    add_bullet(doc, "Create an **Account** and related **Address**")
    add_bullet(doc, "Create a **Contract** and **Contract Pricing**")
    add_bullet(doc, "Create a **Quote**, add products in **STLE**, and generate documents")
    add_bullet(doc, "Submit the quote for **Approval** when required")
    add_bullet(doc, "Create and activate an **Order** (integrates to SyteLine)")

    # ----- 2. Account -----
    add_heading_styled(doc, "2. Create an Account", level=1)
    add_heading_styled(doc, "2.1 Prerequisites", level=2)
    add_bullet(doc, "Access to the **Vulcan Industries** (or **Vulcan Industrial**) application")

    add_heading_styled(doc, "2.2 Steps", level=2)
    add_numbered(doc, "Open the **Vulcan Industries** application from the App Launcher.")
    add_numbered(doc, "Navigate to **Accounts** and click **New**.")
    add_numbered(doc, "Complete the required fields:")
    add_bullet(doc, "**Account Name**", level=1)
    add_bullet(doc, "**Customer Id**", level=1)
    add_bullet(doc, "**Billing Address**", level=1)
    add_numbered(doc, "Click **Save**.")

    add_figure(
        doc,
        "image1.png",
        "Figure 1. New Account dialog — complete required Account Information, then Save.",
        width=5.8,
    )

    # ----- 3. Address -----
    add_heading_styled(doc, "3. Create an Address for the Account", level=1)
    add_body(
        doc,
        "Addresses are managed from the Account related lists. The Parent field on "
        "an Address record is a lookup to a Location record.",
    )

    add_heading_styled(doc, "3.1 Steps", level=2)
    add_numbered(doc, "Open the Account you created.")
    add_numbered(doc, "In **Related List Quick Links**, select **Addresses**.")
    add_numbered(doc, "Click **New**.")
    add_numbered(doc, "Complete the required fields:")
    add_bullet(
        doc,
        "**Parent** — lookup to a **Location**. If needed, create a Location first "
        "(**+ New Location**), save it, then select it.",
        level=1,
    )
    add_bullet(doc, "**Location Type**", level=1)
    add_numbered(doc, "Save the Address.")

    add_figure(
        doc,
        "image2.png",
        "Figure 2. From the Account record, open Addresses in Related List Quick Links.",
        width=5.8,
    )
    add_figure(
        doc,
        "image3.png",
        "Figure 3. New Address — Parent is a Location lookup; create a Location if one does not exist.",
        width=5.8,
    )

    add_callout(
        doc,
        "TIP",
        "Parent = Location. Create and save the Location before completing the Address, "
        "or use + New Location from the Parent lookup.",
        bg=TIP_BG,
    )

    # ----- 4. Contract -----
    add_heading_styled(doc, "4. Contract and Contract Pricing", level=1)
    add_body(
        doc,
        "Contracted list prices appear in the Sales Transaction Line Editor (STLE) "
        "only when the contract is Activated and Contract Pricing Entries are refreshed.",
    )

    add_heading_styled(doc, "4.1 Create a Contract", level=2)
    add_numbered(doc, "Open the **Contracts** object.")
    add_numbered(doc, "Click **New** and complete the required Contract fields.")
    add_numbered(doc, "Click **Save**.")

    add_heading_styled(doc, "4.2 Create Contract Line / Contract Item Pricing", level=2)
    add_numbered(
        doc,
        "On the Contract record, open the related list for contract line / pricing "
        "(for example, **Contract Item Prices**).",
    )
    add_numbered(doc, "Click **New**.")
    add_numbered(
        doc,
        "Enter pricing details (for example: Item, Contract, Price, Start Date/Time, "
        "and Product Selling Model).",
    )
    add_numbered(doc, "Click **Save**.")

    add_figure(
        doc,
        "image4.png",
        "Figure 4. New Contract Item Price — set Item, Contract, Price, and Start Date.",
        width=5.6,
    )

    add_heading_styled(doc, "4.3 Activate the Contract", level=2)
    add_numbered(doc, "Activate the Contract.")
    add_callout(
        doc,
        "IMPORTANT",
        "Only an activated contract drives Contracted List Price in STLE. "
        "Inactive contracts will not reflect contracted pricing on quote lines.",
        bg=IMPORTANT_BG,
    )

    add_heading_styled(doc, "4.4 Refresh Contract Pricing Entries", level=2)
    add_numbered(
        doc,
        "Refresh the **Contract Pricing Entries** decision table so pricing is available to quotes.",
    )

    add_heading_styled(doc, "4.5 Create a Quote from the Contract (optional path)", level=2)
    add_numbered(doc, "On the Contract related list, click **New Quote**.")
    add_callout(
        doc,
        "NOTE",
        "When a quote is created from a Contract, the quote stores Contract Id. "
        "That Contract Id is what allows list price changes and contracted pricing to appear in STLE.",
        bg=CALLOUT_BG,
    )

    # ----- 5. Quote from Account -----
    add_heading_styled(doc, "5. Create a Quote (from Account)", level=1)

    add_heading_styled(doc, "5.1 Steps", level=2)
    add_numbered(doc, "Open the Account record.")
    add_numbered(doc, "Click **Create Quote** on the record page.")
    add_numbered(doc, "Enter **Quote Name** and **Start Date**, then save.")
    add_numbered(doc, "On the Quote record, complete the following fields:")
    add_bullet(doc, "**Billing Term**", level=1)
    add_bullet(doc, "**Shipping Term** (Ship Term)", level=1)
    add_bullet(doc, "**Warehouse**", level=1)
    add_bullet(doc, "**Address**", level=1)

    add_figure(
        doc,
        "image5.png",
        "Figure 5. Account action bar — click Create Quote.",
        width=6.0,
    )

    add_callout(
        doc,
        "AUTOMATION",
        "The system can populate the latest Contract Id on the Quote when contract "
        "pricing is available for the Quote’s Account.",
        bg=CALLOUT_BG,
    )

    # ----- 6. STLE -----
    add_heading_styled(doc, "6. Add Products in STLE", level=1)
    add_body(
        doc,
        "Use the Sales Transaction Line Editor (STLE) on the Quote Lines tab to browse "
        "the catalog and add products.",
    )

    add_heading_styled(doc, "6.1 Steps", level=2)
    add_numbered(doc, "On the Quote, open the **Lines** tab.")
    add_numbered(doc, "In STLE, click **Browse Catalogs** (or use **Browse Catalog** / **Add Product**).")
    add_numbered(doc, "Select a catalog and add products by category into STLE.")
    add_numbered(doc, "Review list price and contracted pricing on the lines as applicable.")

    add_figure(
        doc,
        "image6.png",
        "Figure 6. Quote Lines tab — add products via catalog / Add Product in STLE.",
        width=6.0,
    )

    # ----- 7. Quote documents -----
    add_heading_styled(doc, "7. Generate and Send the Quote Document", level=1)

    add_heading_styled(doc, "7.1 Steps", level=2)
    add_numbered(
        doc,
        "On the Quote record, click **Generate Quote PDF** / **Generate Quote Document** "
        "to preview and save the document to **Files**.",
    )
    add_numbered(doc, "Click **Send Quote Doc** to email the file to a user.")

    add_figure(
        doc,
        "image7.png",
        "Figure 7. Quote header — Generate Quote PDF and Send Quote Doc.",
        width=6.0,
    )

    # ----- 8. Approvals -----
    add_heading_styled(doc, "8. Quote Approvals", level=1)

    add_heading_styled(doc, "8.1 Business rule", level=2)
    add_body(
        doc,
        "If the quote margin is less than 10%, clicking Submit for Approval routes "
        "the quote to a VP user for approval.",
    )
    add_callout(
        doc,
        "STATUS RULE",
        "Approval Status should be updated to Needs Approval when the approval "
        "condition is met (for example, margin below the approved threshold).",
        bg=WARN_BG,
    )

    add_heading_styled(doc, "8.2 Behavior after submit", level=2)
    add_bullet(doc, "Once submitted for approval, the quote status changes to **In Review**.")
    add_bullet(doc, "Approvers receive an email notification.")
    add_bullet(
        doc,
        "After submit, the Quote **Related** list / **Approval History** provides "
        "options to **Approve**, **Reject**, or **Recall**.",
    )

    add_heading_styled(doc, "8.3 Quote creation restriction", level=2)
    add_body(
        doc,
        "Quote creation is allowed only when Approval Status is not equal to any of the following:",
    )
    add_bullet(doc, "**Needs Approval**")
    add_bullet(doc, "**Rejected**")
    add_bullet(doc, "**In Review**")

    add_heading_styled(doc, "8.4 Submit for approval", level=2)
    add_numbered(doc, "On the Quote, review **Margin** and **Approval Status**.")
    add_numbered(doc, "Click **Submit for Approval**.")
    add_numbered(
        doc,
        "On the Quote **Related** lists, open **Approval History** to track status "
        "(Submitted, Approved, Rejected, Recalled, and so on).",
    )

    add_figure(
        doc,
        "image8.png",
        "Figure 8. Submit for Approval, Approval Status, and Margin on the Quote.",
        width=5.4,
    )
    add_figure(
        doc,
        "image9.png",
        "Figure 9. Approval History related list — track step status and assignee.",
        width=5.6,
    )

    add_callout(
        doc,
        "REMINDER",
        "Margin less than 10% requires VP approval. After submit, status is In Review "
        "and an approval email is sent.",
        bg=WARN_BG,
    )

    add_heading_styled(doc, "8.5 Approve, Reject, or Recall", level=2)
    add_body(
        doc,
        "After a Quote is submitted for approval, use Approval History on the Quote "
        "related list to take action:",
    )
    add_bullet(doc, "**Approve** — approve the pending approval request.")
    add_bullet(doc, "**Reject** — reject the pending approval request.")
    add_bullet(
        doc,
        "**Recall** — open the dropdown next to Reject and select **Recall** to "
        "withdraw the approval request.",
    )
    add_callout(
        doc,
        "RECALL STATUS",
        "If approval is recalled, Approval Status returns to Needs Approval.",
        bg=IMPORTANT_BG,
    )

    add_figure(
        doc,
        "image12.png",
        "Figure 10. Approval History — Approve, Reject, and Recall actions on a submitted Quote.",
        width=6.0,
    )

    add_heading_styled(doc, "8.6 Approve or Reject from email", level=2)
    add_body(
        doc,
        "Once a Quote is submitted for approval, Salesforce sends an email notification "
        "to the approver. Approve and Reject options are enabled directly in the email, "
        "so the approver can take action without opening Salesforce.",
    )
    add_numbered(
        doc,
        "The approver receives an email when a Quote requires approval "
        "(for example, when margin is below the approved threshold).",
    )
    add_numbered(
        doc,
        "From the email, the approver can **Approve** or **Reject** "
        "(including by reply action where enabled).",
    )
    add_numbered(
        doc,
        "After the decision, the submitter / team receives the appropriate email "
        "notification that the Quote was approved or rejected.",
    )

    add_figure(
        doc,
        "image14.png",
        "Figure 11. Approval request email — Quote requires approval when margin is below threshold.",
        width=5.2,
    )
    add_figure(
        doc,
        "image13.png",
        "Figure 12. Email approval reply — approver can Approve (or Reject) from email.",
        width=4.8,
    )

    add_heading_styled(doc, "8.7 Approval status summary", level=2)
    add_bullet(
        doc,
        "When the approval condition is met → Approval Status = **Needs Approval**.",
    )
    add_bullet(doc, "After Submit for Approval → status = **In Review**.")
    add_bullet(doc, "If approval is **Recalled** → status returns to **Needs Approval**.")
    add_bullet(doc, "If approved or rejected → the related party receives email notification.")

    # ----- 9. Create Order -----
    add_heading_styled(doc, "9. Create an Order from a Quote", level=1)

    add_heading_styled(doc, "9.1 Steps", level=2)
    add_numbered(doc, "On the Quote, click **Create Order**.")
    add_numbered(doc, "Select **Create Single Order**, then finish.")
    add_numbered(doc, "Confirm the Order is linked to the correct **Account** and **Contract**.")

    add_figure(
        doc,
        "image10.png",
        "Figure 13. Create Order — select Create Single Order, then continue.",
        width=5.6,
    )

    add_callout(
        doc,
        "NOTE",
        "Each order is associated with the Contract and Account record.",
        bg=CALLOUT_BG,
    )

    # ----- 10. Order docs & activation -----
    add_heading_styled(doc, "10. Order Documents and Activation", level=1)

    add_heading_styled(doc, "10.1 Generate and send order documents", level=2)
    add_numbered(doc, "Open the Order.")
    add_numbered(
        doc,
        "Click **Generate Order Doc** to view or download the document "
        "(saved to related **Files**).",
    )
    add_numbered(doc, "Click **Send Order Doc** to email the attached document.")

    add_figure(
        doc,
        "image11.png",
        "Figure 14. Order header — Generate Order Doc and Send Order Doc.",
        width=6.0,
    )

    add_heading_styled(doc, "10.2 Activate the Order", level=2)
    add_numbered(doc, "Open the Order and **Activate** it.")

    add_callout(
        doc,
        "INTEGRATION",
        "Activated Orders integrate and create orders in SyteLine for the shipping team. "
        "There are no updates back into Salesforce from SyteLine after activation.",
        bg=IMPORTANT_BG,
    )

    # ----- 11. Quick reference -----
    add_heading_styled(doc, "11. Quick Reference — Key Tips & Rules", level=1)
    add_body(
        doc,
        "Use this table as a quick checklist. All guidance from earlier sections is summarized here.",
    )

    ref_rows = [
        ("Area", "Guidance"),
        ("Account", "Required: Account Name, Customer Id, Billing Address"),
        ("Address Parent", "Lookup to Location — create Location if missing"),
        ("Contract", "Must be Activated for contracted list price in STLE"),
        ("Pricing refresh", "Refresh Contract Pricing Entries decision table after pricing changes"),
        (
            "Quote ← Contract",
            "Quote carries Contract Id; required for contracted pricing in STLE",
        ),
        (
            "Quote ← Account",
            "Automation may set latest Contract Id when pricing exists for the Account",
        ),
        ("Quote fields", "Billing Term, Shipping Term, Warehouse, Address"),
        ("STLE", "Lines tab → Browse Catalog / Add Product"),
        ("Approval", "Margin < 10% → Submit for Approval → VP; status In Review"),
        (
            "Approve / Reject / Recall",
            "From Approval History after submit; Recall returns status to Needs Approval",
        ),
        (
            "Email approval",
            "Approver can Approve or Reject from email; notifications sent on decision",
        ),
        (
            "Approval statuses",
            "Needs Approval (condition met / after recall) → In Review (submitted)",
        ),
        (
            "Quote gate",
            "Do not create when Approval Status is Needs Approval, Rejected, or In Review",
        ),
        ("Order", "Create Single Order from Quote"),
        ("Activation", "Activated Order → SyteLine; no inbound updates to Salesforce"),
    ]

    table = doc.add_table(rows=len(ref_rows), cols=2)
    table.style = "Table Grid"
    for i, (a, b) in enumerate(ref_rows):
        c0, c1 = table.cell(i, 0), table.cell(i, 1)
        c0.text = ""
        c1.text = ""
        r0 = c0.paragraphs[0].add_run(a)
        r1 = c1.paragraphs[0].add_run(b)
        if i == 0:
            set_run_font(r0, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            set_run_font(r1, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            set_cell_shading(c0, "017EC5")
            set_cell_shading(c1, "017EC5")
        else:
            set_run_font(r0, size=10, bold=True)
            set_run_font(r1, size=10)
            if i % 2 == 0:
                set_cell_shading(c0, "F4F6F9")
                set_cell_shading(c1, "F4F6F9")

    footer = doc.add_paragraph()
    set_paragraph_spacing(footer, before=24, after=0)
    add_horizontal_line(footer)
    end = doc.add_paragraph()
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = end.add_run("— End of User Guide —")
    set_run_font(run, size=10, color=GRAY)
    run.italic = True

    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    build()
